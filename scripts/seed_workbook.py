#!/usr/bin/env python3
"""One-time seed of `S1 NY Islanders.xlsx` from the franchise screens captured on
09/29/2026 (in-game date). Run once; after that the workbook is the data source and
this script is only a record of where the opening data came from.

Every sheet: row 1 headers, row 2 blank styled template, data from row 3.
"""
import sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "S1 NY Islanders.xlsx"

HEAD_FILL = PatternFill("solid", fgColor="00539B")
HEAD_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="D7DFEB")
BORDER = Border(bottom=THIN)

def sheet(wb, name, headers, rows, widths=None):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEAD_FILL; cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    # row 2: blank styled template that add_game.py copies formatting from
    for c in range(1, len(headers) + 1):
        ws.cell(row=2, column=c).border = BORDER
    for r in rows:
        ws.append(list(r))
    for c, h in enumerate(headers, start=1):
        w = (widths or {}).get(h)
        if w is None:
            w = max([len(str(h))] + [len(str(r[c - 1])) for r in rows if c - 1 < len(r) and r[c - 1] is not None]) + 2
        ws.column_dimensions[get_column_letter(c)].width = min(max(w, 6), 44)
    ws.freeze_panes = "A3"
    return ws

# --------------------------------------------------------------------------- Roster Ref
# Player, Pos, Group, Status, #, OVR, POT, POT Cert, Age, Ht, Wt, Shoots, Type, Ext, Clause, FSC,
# 26-27 .. 33-34 (salary in $M; RFA/UFA/UNSIGNED where the game shows a tag), Then
ROSTER_HEADERS = ["Player", "Pos", "Group", "Status", "#", "OVR", "POT", "POT Cert", "Age",
                  "Ht", "Wt", "Shoots", "Type", "Ext", "Clause", "FSC",
                  "26-27", "27-28", "28-29", "29-30", "30-31", "31-32", "32-33", "33-34", "Then"]

def contract(first, through, then):
    """Flat salary `first` ($M) for every season up to and including `through`, then `then`."""
    seasons = ["26-27", "27-28", "28-29", "29-30", "30-31", "31-32", "32-33", "33-34"]
    out = []
    idx = seasons.index(through)
    for i, s in enumerate(seasons):
        if i <= idx:
            out.append(first)
        elif i == idx + 1:
            out.append(then)
        else:
            out.append(None)
    return out

M, S = "Main Roster", "In the System"
D, SCR = "Dressed", "Scratched"
# name, pos, group, status, num, ovr, pot, cert, age, ht, wt, shoots, type, ext, clause, fsc, salary, through, then
SK = [
 ("M. Schaefer",  "LD/RD", M, D,   48, 91, "Elite",      "High",  19, None,  None, None, None, "-",   "-",     "No",  0.975, "27-28", "RFA"),
 ("M. Barzal",    "C/RW",  M, D,   None, 88, "Elite",    "Exact", 29, None,  None, None, None, "-",   "M-NTC", "Yes", 9.150, "30-31", "UFA"),
 ("B. Horvat",    "C",     M, D,   14, 88, "Elite",      "Exact", 31, "6'1\"", 225, None, "TWF", "-", "NTC",   "Yes", 8.500, "30-31", "UFA"),
 ("R. Pulock",    "RD",    M, D,   None, 88, "Elite",    "Exact", 31, None,  None, None, None, "-",   "NTC",   "Yes", 6.150, "29-30", "UFA"),
 ("A. Pelech",    "LD",    M, D,   None, 87, "Top 4 D",  "Exact", 32, None,  None, None, None, "-",   "M-NTC", "Yes", 5.750, "28-29", "UFA"),
 ("M. Weegar",    "LD/RD", M, D,   None, 86, "Top 6 D",  "Exact", 32, None,  None, None, None, "-",   "NTC",   "Yes", 6.250, "30-31", "UFA"),
 ("M. Coronato",  "RW/LW", M, D,   27, 85, "Top 6 F",    "Med",   23, "5'10\"", 183, "R", "SNP", "-", "-",     "Yes", 6.500, "31-32", "UFA"),
 ("S. Holmstrom", "RW",    M, D,   None, 84, "Top 9 F",  "Med",   25, None,  None, None, None, "No",  "-",     "No",  3.625, "26-27", "RFA"),
 ("B. Schenn",    "C/LW",  M, D,   None, 84, "Top 6 F",  "Exact", 35, None,  None, None, None, "-",   "M-NTC", "Yes", 6.500, "27-28", "UFA"),
 ("T. DeAngelo",  "RD",    M, SCR, None, 83, "Top 4 D",  "Exact", 30, None,  None, None, None, "-",   "M-NTC", "Yes", 4.500, "27-28", "UFA"),
 ("M. Maccelli",  "LW",    M, D,   None, 83, "Top 6 F",  "Low",   25, None,  None, None, None, "-",   "-",     "No",  2.250, "26-27", "RFA"),
 ("A. Romanov",   "LD",    M, D,   None, 83, "Top 4 D",  "Med",   26, None,  None, None, None, "-",   "NTC",   "Yes", 6.250, "32-33", "UFA"),
 ("E. Heineman",  "LW",    M, D,   51, 82, "Top 6 F",    "Med",   24, "6'2\"", 204, "L", "PWF", "Yes", "-",    "No",  1.100, "26-27", "RFA"),
 ("K. Palmieri",  "RW/LW", M, D,   21, 82, "Top 6 F",    "Exact", 35, "6'0\"", 192, None, "TWF", "Yes", "M-NTC", "No", 4.750, "26-27", "UFA"),
 ("C. Cizikas",   "C",     M, D,   None, 81, "Top 9 F",  "Exact", 35, None,  None, None, None, "Yes", "-",     "No",  2.500, "26-27", "UFA"),
 ("A. Duclair",   "LW/RW", M, D,   None, 81, "Top 6 F",  "Exact", 31, None,  None, None, None, "-",   "M-NTC", "Yes", 3.500, "27-28", "UFA"),
 ("J. Pageau",    "C",     M, D,   44, 81, "Top 6 F",    "Exact", 33, "5'11\"", 180, None, "TWF", "-", "NTC",  "Yes", 4.850, "28-29", "UFA"),
 ("C. Ritchie",   "C",     M, SCR, 64, 81, "Top 6 F",    "Med",   21, "6'2\"", 200, None, None, "-",  "-",     "No",  0.950, "26-27", "RFA"),
 ("P. Engvall",   "LW/RW", M, D,   None, 80, "Bottom 6 F","Exact", 30, None,  None, None, None, "-",  "M-NTC", "Yes", 3.000, "29-30", "UFA"),
 ("M. Kessel",    "RD/LD", M, D,   None, 79, "Top 6 D",  "Med",   26, None,  None, None, None, "-",   "-",     "No",  0.850, "26-27", "UFA"),
 ("K. MacLean",   "LW",    M, SCR, None, 79, "Bottom 6 F","Exact", 27, None,  None, None, None, "-",  "-",     "No",  0.850, "26-27", "UFA"),
 # in the system
 ("M. Chaffee",   "RW",    S, None, None, 77, "Bottom 6 F","Exact", 28, None, None, None, None, "-",  "-",     "No",  0.850, "26-27", "UFA"),
 ("I. George",    "LD",    S, None, None, 77, "Top 6 D",  "Med",   22, None,  None, None, None, "Yes", "-",    "No",  0.915, "26-27", "RFA"),
 ("E. Bear",      "RD",    S, None, None, 76, "Top 4 D",  "Exact", 29, None,  None, None, None, "-",   "-",     "No",  0.850, "26-27", "UFA"),
 ("V. Eklund",    "LW",    S, None, None, 75, "Elite",    "Med",   19, None,  None, None, None, "-",   "-",     "No",  0.975, "27-28", "RFA"),
 ("L. Foudy",     "C/LW",  S, None, None, 75, "Top 6 F",  "Low",   26, None,  None, None, None, "-",   "-",     "No",  0.850, "26-27", "UFA"),
 ("M. Warren",    "LD",    S, None, None, 74, "Top 6 D",  "Med",   25, None,  None, None, None, "-",   "-",     "No",  0.850, "26-27", "RFA"),
 ("M. Luff",      "RW",    S, None, None, 73, "Bottom 6 F","Exact", 29, None, None, None, None, "-",  "-",     "No",  0.850, "26-27", "UFA"),
 ("K. Aitcheson", "LD/RD", S, None, None, 72, "Top 4 D",  "Med",   20, None,  None, None, None, "-",   "-",     "No",  1.075, "28-29", "RFA"),
 ("A. Jefferies", "LW",    S, None, None, 71, "Top 9 F",  "Med",   24, None,  None, None, None, "-",   "-",     "No",  0.850, "26-27", "RFA"),
 ("D. Kuefler",   "LW",    S, None, None, 70, "Bottom 6 F","Med",   24, None, None, None, None, "-",  "-",     "No",  0.875, "27-28", "RFA"),
 ("M. Gustafsson","LD",    S, None, None, 69, "Top 4 D",  "Med",   18, None,  None, None, None, "-",   "-",     "No",  1.075, "28-29", "RFA"),
 ("E. Liukas",    "RW",    S, None, None, 68, "Bottom 6 F","Med",   24, None, None, None, None, "-",  "-",     "No",  0.850, "26-27", "RFA"),
 ("C. Odelius",   "LD/RD", S, None, None, 68, "Top 6 D",  "Med",   22, None,  None, None, None, "Yes", "-",    "No",  0.865, "26-27", "RFA"),
 ("J. Larson",    "RW/LW", S, None, None, 67, "Bottom 6 F","Med",   25, None, None, None, None, "-",  "-",     "No",  0.875, "27-28", "UFA"),
 ("J. Pulkkinen", "LD/RD", S, None, None, 67, "Top 6 D",  "Med",   21, None,  None, None, None, "-",   "-",     "No",  0.880, "26-27", "RFA"),
 ("L. Romano",    "RW",    S, None, None, 66, "Bottom 6 F","Med",   19, None, None, None, None, "-",  "-",     "No",  "UNSIGNED", "26-27", "UFA"),
 ("G. Veremyev",  "C/LW",  S, None, None, 64, "Bottom 6 F","Med",   23, None, None, None, None, "-",  "-",     "No",  0.885, "26-27", "RFA"),
 ("J. Kvasnicka", "RW",    S, None, None, 62, "Bottom 6 F","Med",   19, None, None, None, None, "-",  "-",     "No",  "UNSIGNED", "26-27", "UFA"),
 ("T. Poletin",   "LW",    S, None, None, 61, "Bottom 6 F","Med",   19, None, None, None, None, "-",  "-",     "No",  "UNSIGNED", "26-27", "UFA"),
 ("V. Dravecky",  "RD/LD", S, None, None, 60, "Top 6 D",  "Low",   18, None,  None, None, None, "-",   "-",     "No",  "UNSIGNED", "27-28", "UFA"),
 ("J. Nurmi",     "LW",    S, None, None, 58, "Top 9 F",  "Med",   21, None,  None, None, None, "Yes", "-",    "No",  0.870, "26-27", "RFA"),
]
GOALIES = [  # contracts pending from the user
 ("I. Sorokin",   "G", M, D, None, 93, None, None, None, None, None, None, None, None, None, None, None, None, None),
 ("S. Varlamov",  "G", M, D, 40,   81, None, None, None, "6'2\"", 201, None, None, None, None, None, None, None, None),
]

def roster_rows():
    rows = []
    for p in SK:
        (name, pos, grp, st, num, ovr, pot, cert, age, ht, wt, sh, typ, ext, cl, fsc, sal, thr, then) = p
        rows.append([name, pos, grp, st, num, ovr, pot, cert, age, ht, wt, sh, typ, ext, cl, fsc] + contract(sal, thr, then) + [then])
    for g in GOALIES:
        (name, pos, grp, st, num, ovr, pot, cert, age, ht, wt, sh, typ, ext, cl, fsc, sal, thr, then) = g
        rows.append([name, pos, grp, st, num, ovr, pot, cert, age, ht, wt, sh, typ, ext, cl, fsc] + [None] * 8 + [None])
    return rows

# --------------------------------------------------------------------------- Lines
# Unit, Slot, Player, Pos, OVR, Chem  (chem is the unit's line-chemistry impact, repeated per row)
LINES = [
 ("F1", "LW", "B. Horvat", "LW", 88, 0), ("F1", "C", "M. Barzal", "C", 88, 0), ("F1", "RW", "K. Palmieri", "RW", 82, 0),
 ("F2", "LW", "M. Maccelli", "LW", 83, 2), ("F2", "C", "B. Schenn", "C", 84, 2), ("F2", "RW", "P. Engvall", "RW", 80, 2),
 ("F3", "LW", "E. Heineman", "LW", 82, 1), ("F3", "C", "C. Cizikas", "C", 81, 1), ("F3", "RW", "M. Coronato", "RW", 85, 1),
 ("F4", "LW", "A. Duclair", "LW", 81, 0), ("F4", "C", "J. Pageau", "C", 81, 0), ("F4", "RW", "S. Holmstrom", "RW", 84, 0),
 ("D1", "LD", "M. Schaefer", "LD", 91, 0), ("D1", "RD", "R. Pulock", "RD", 88, 0),
 ("D2", "LD", "A. Pelech", "LD", 87, 4), ("D2", "RD", "M. Weegar", "RD", 86, 4),
 ("D3", "LD", "A. Romanov", "LD", 83, 0), ("D3", "RD", "M. Kessel", "RD", 79, 0),
 ("PP1", "LW", "B. Horvat", "LW", 88, 5), ("PP1", "C", "M. Barzal", "C", 88, 5), ("PP1", "RW", "S. Holmstrom", "RW", 84, 5),
 ("PP1", "RD", "M. Schaefer", "RD", 91, 5), ("PP1", "LD", "M. Coronato", "LD", 85, 5),
 ("PP2", "LW", "E. Heineman", "LW", 82, 0), ("PP2", "C", "B. Schenn", "C", 84, 0), ("PP2", "RW", "M. Maccelli", "RW", 83, 0),
 ("PP2", "RD", "R. Pulock", "RD", 88, 0), ("PP2", "LD", "A. Duclair", "LD", 81, 0),
 ("PK1", "C", "C. Cizikas", "C", 81, 4), ("PK1", "LW", "J. Pageau", "LW", 81, 4), ("PK1", "RD", "A. Pelech", "RD", 87, 4), ("PK1", "LD", "M. Weegar", "LD", 86, 4),
 ("PK2", "C", "B. Schenn", "C", 84, 1), ("PK2", "LW", "B. Horvat", "LW", 88, 1), ("PK2", "RD", "R. Pulock", "RD", 88, 1), ("PK2", "LD", "A. Romanov", "LD", 83, 1),
 ("PK3", "C", "M. Coronato", "C", 85, 1), ("PK3", "LW", "K. Palmieri", "LW", 82, 1), ("PK3", "RD", "M. Kessel", "RD", 79, 1), ("PK3", "LD", "M. Schaefer", "LD", 91, 1),
 ("G", "1", "S. Varlamov", "G", 81, None), ("G", "2", "I. Sorokin", "G", 93, None),
 ("SCR", "", "C. Ritchie", "C", 81, None), ("SCR", "", "T. DeAngelo", "RD", 83, None), ("SCR", "", "K. MacLean", "LW", 79, None),
]

OWNER_GOALS = [
 ("Primary",   "Improve the Top 6 forward group",            50000, "03/01/2027", "Open"),
 ("Secondary", "Win the regular season home opener",         30000, "10/03/2026", "Open"),
 ("Secondary", "43 wins this season",                        40000, "04/11/2027", "Open"),
 ("Stretch",   "Upgrade the parking lot by 1 level",         30000, "06/19/2027", "Open"),
]

BUDGET = [
 ("Player Salaries", 92.685, 0.000, 92.685),
 ("Arena Operations", 2.302, 0.000, 2.302),
 ("Promotions",       0.977, 0.000, 0.977),
 ("Advertising",      1.738, 0.003, 1.735),
 ("Scout Salaries",   3.101, 0.000, 3.101),
 ("Scout Travel",     2.029, 0.002, 2.027),
 ("Coach Budget",     8.915, 0.000, 8.915),
]

CAP = [  # Season, Salary Cap, Main Roster ($M, skaters), System ($M, skaters), Contracts -- as the game shows them
 ("26-27", 104.000, 88.750, 15.240, 43),
 ("27-28", None,    71.875, 4.875,  None),
 ("28-29", None,    56.400, 2.150,  None),
 ("29-30", 106.600, 45.800, 0.000,  8),
 ("30-31", 109.600, 36.650, 0.000,  6),
 ("31-32", None,    12.750, 0.000,  None),
 ("32-33", 111.100, 6.250,  0.000,  1),
 ("33-34", 112.600, 0.000,  0.000,  0),
]

FRONT = [
 ("As Of",                    "09/29/2026"),
 ("Owner Happiness",          "High"),
 ("Owner Message",            "I'm very pleased with the work you have been doing with this franchise. Keep up the good work!"),
 ("State of the Team",        "Buyer"),
 ("Funds Remaining ($M)",     0.559),
 ("Salary Target ($M)",       63.885),
 ("Salary Cap ($M)",          104.000),
 ("Salary Cap Floor ($M)",    76.900),
 ("Max Player Salary ($M)",   20.800),
 ("Min Player Salary ($M)",   0.850),
 ("Max Rookie Salary ($M)",   1.025),
 ("Team Cap Hit ($M)",        99.750),
 ("Cap Space ($M)",           4.250),
 ("Playoff Cap ($M)",         93.450),
 ("Contracts",                "43/50"),
 ("Exempt Contracts",         "0/40"),
 ("Retained Salary ($M)",     0),
 ("Buyout Penalty ($M)",      0),
]

TRADES = [
 ("09/29/2026", "TBL", "Incoming", "NYI 2027 R3; I. George (D, $0.915M)", "TOR 2027 R4; TBL 2027 R4", "Declined"),
 ("09/29/2026", "STL", "Incoming", "COL 2028 R3; J. Nurmi (LW, $0.870M)", "STL 2028 R4; STL 2028 R7", "Declined"),
]

SCHEDULE = [  # G, Date (mm/dd/yyyy), H/A, Opp, Time (ET) -- NHL.com official 2026-27 release, 84 games
 (1, "09/30/2026", "A", "TOR", "7:30p"),
 (2, "10/03/2026", "H", "NJD", "7:30p"),
 (3, "10/06/2026", "A", "NYR", "7p"),
 (4, "10/08/2026", "H", "CHI", "7:30p"),
 (5, "10/10/2026", "H", "TBL", "7:30p"),
 (6, "10/13/2026", "H", "VAN", "7:45p"),
 (7, "10/15/2026", "A", "OTT", "7p"),
 (8, "10/17/2026", "A", "TOR", "7p"),
 (9, "10/20/2026", "H", "ANA", "7p"),
 (10, "10/22/2026", "H", "LAK", "7:30p"),
 (11, "10/25/2026", "A", "DAL", "6p"),
 (12, "10/27/2026", "A", "NSH", "8p"),
 (13, "10/29/2026", "A", "MIN", "8p"),
 (14, "10/31/2026", "H", "EDM", "3:30p"),
 (15, "11/05/2026", "H", "CAR", "7p"),
 (16, "11/07/2026", "H", "NJD", "7p"),
 (17, "11/09/2026", "A", "SJS", "10p"),
 (18, "11/12/2026", "A", "LAK", "10p"),
 (19, "11/14/2026", "A", "ANA", "10p"),
 (20, "11/17/2026", "H", "CBJ", "7p"),
 (21, "11/19/2026", "A", "WPG", "8p"),
 (22, "11/21/2026", "A", "DET", "7p"),
 (23, "11/23/2026", "H", "TOR", "7:30p"),
 (24, "11/25/2026", "H", "STL", "7p"),
 (25, "11/27/2026", "H", "WPG", "7:30p"),
 (26, "11/28/2026", "A", "PHI", "7:30p"),
 (27, "12/01/2026", "H", "FLA", "7p"),
 (28, "12/03/2026", "A", "PHI", "7p"),
 (29, "12/04/2026", "H", "SJS", "7p"),
 (30, "12/07/2026", "H", "COL", "1p"),
 (31, "12/10/2026", "H", "CGY", "7p"),
 (32, "12/11/2026", "A", "CAR", "7p"),
 (33, "12/15/2026", "A", "UTA", "9p"),
 (34, "12/16/2026", "A", "COL", "9p"),
 (35, "12/18/2026", "A", "VGK", "10p"),
 (36, "12/20/2026", "H", "NYR", "7p"),
 (37, "12/22/2026", "H", "BOS", "7p"),
 (38, "12/26/2026", "A", "BOS", "7p"),
 (39, "12/27/2026", "H", "OTT", "7p"),
 (40, "12/30/2026", "H", "WSH", "4p"),
 (41, "01/02/2027", "A", "SEA", "10p"),
 (42, "01/05/2027", "A", "VAN", "9p"),
 (43, "01/07/2027", "A", "CGY", "8p"),
 (44, "01/09/2027", "A", "EDM", "3:30p"),
 (45, "01/13/2027", "H", "PIT", "7:30p"),
 (46, "01/15/2027", "H", "DAL", "7p"),
 (47, "01/16/2027", "A", "NJD", "7p"),
 (48, "01/18/2027", "H", "NSH", "3p"),
 (49, "01/20/2027", "H", "BUF", "7p"),
 (50, "01/22/2027", "A", "WSH", "7p"),
 (51, "01/23/2027", "H", "SEA", "7:30p"),
 (52, "01/26/2027", "A", "PIT", "7p"),
 (53, "01/28/2027", "A", "CAR", "7p"),
 (54, "01/30/2027", "A", "TBL", "7p"),
 (55, "02/01/2027", "A", "FLA", "1p"),
 (56, "02/03/2027", "A", "BUF", "7p"),
 (57, "02/12/2027", "A", "NYR", "7p"),
 (58, "02/13/2027", "H", "MIN", "7p"),
 (59, "02/15/2027", "H", "UTA", "3p"),
 (60, "02/18/2027", "H", "NYR", "7p"),
 (61, "02/20/2027", "A", "MTL", "7p"),
 (62, "02/22/2027", "A", "CBJ", "7p"),
 (63, "02/23/2027", "A", "CHI", "8p"),
 (64, "02/26/2027", "H", "PHI", "7p"),
 (65, "02/27/2027", "A", "CBJ", "7p"),
 (66, "03/01/2027", "H", "WSH", "7p"),
 (67, "03/04/2027", "H", "VGK", "7p"),
 (68, "03/07/2027", "H", "PIT", "1p"),
 (69, "03/09/2027", "H", "MTL", "7p"),
 (70, "03/12/2027", "A", "FLA", "7p"),
 (71, "03/13/2027", "A", "TBL", "7p"),
 (72, "03/15/2027", "A", "BOS", "7p"),
 (73, "03/17/2027", "A", "WSH", "7p"),
 (74, "03/20/2027", "H", "PHI", "3p"),
 (75, "03/21/2027", "H", "BUF", "3p"),
 (76, "03/25/2027", "A", "STL", "8p"),
 (77, "03/27/2027", "H", "DET", "7:30p"),
 (78, "03/30/2027", "H", "CAR", "7p"),
 (79, "04/01/2027", "A", "PIT", "7p"),
 (80, "04/02/2027", "H", "OTT", "7p"),
 (81, "04/04/2027", "H", "MTL", "7:30p"),
 (82, "04/06/2027", "H", "DET", "7:30p"),
 (83, "04/08/2027", "A", "NJD", "7p"),
 (84, "04/10/2027", "H", "CBJ", "5p"),
]

TEAMS = [
 ("Anaheim Ducks","ANA","West","Pacific"),("Boston Bruins","BOS","East","Atlantic"),("Buffalo Sabres","BUF","East","Atlantic"),
 ("Calgary Flames","CGY","West","Pacific"),("Carolina Hurricanes","CAR","East","Metropolitan"),("Chicago Blackhawks","CHI","West","Central"),
 ("Colorado Avalanche","COL","West","Central"),("Columbus Blue Jackets","CBJ","East","Metropolitan"),("Dallas Stars","DAL","West","Central"),
 ("Detroit Red Wings","DET","East","Atlantic"),("Edmonton Oilers","EDM","West","Pacific"),("Florida Panthers","FLA","East","Atlantic"),
 ("Los Angeles Kings","LAK","West","Pacific"),("Minnesota Wild","MIN","West","Central"),("Montreal Canadiens","MTL","East","Atlantic"),
 ("Nashville Predators","NSH","West","Central"),("New Jersey Devils","NJD","East","Metropolitan"),("New York Islanders","NYI","East","Metropolitan"),
 ("New York Rangers","NYR","East","Metropolitan"),("Ottawa Senators","OTT","East","Atlantic"),("Philadelphia Flyers","PHI","East","Metropolitan"),
 ("Pittsburgh Penguins","PIT","East","Metropolitan"),("San Jose Sharks","SJS","West","Pacific"),("Seattle Kraken","SEA","West","Pacific"),
 ("St. Louis Blues","STL","West","Central"),("Tampa Bay Lightning","TBL","East","Atlantic"),("Toronto Maple Leafs","TOR","East","Atlantic"),
 ("Utah Mammoth","UTA","West","Central"),("Vancouver Canucks","VAN","West","Pacific"),("Vegas Golden Knights","VGK","West","Pacific"),
 ("Washington Capitals","WSH","East","Metropolitan"),("Winnipeg Jets","WPG","West","Central"),
]

def main():
    if OUT.exists() and "--force" not in sys.argv:
        sys.exit(f"{OUT.name} exists; pass --force to overwrite")
    wb = Workbook(); wb.remove(wb.active)
    sheet(wb, "Roster Ref", ROSTER_HEADERS, roster_rows(), {"Player": 16})
    sheet(wb, "Lines", ["Unit", "Slot", "Player", "Pos", "OVR", "Chem"], LINES)
    sheet(wb, "Owner Goals", ["Tier", "Goal", "Reward", "Eval Date", "Status"], OWNER_GOALS, {"Goal": 44})
    sheet(wb, "Budget", ["Line", "Allocated", "Spent", "Remaining"], BUDGET)
    sheet(wb, "Cap", ["Season", "Salary Cap", "Main Roster", "System", "Contracts"], CAP)
    sheet(wb, "Front Office", ["Key", "Value"], FRONT, {"Value": 44})
    sheet(wb, "Trades", ["Date", "Partner", "Direction", "Out", "In", "Result"], TRADES, {"Out": 36, "In": 30})
    sheet(wb, "Schedule", ["G", "Date", "H/A", "Opp", "Time (ET)"], SCHEDULE)
    sheet(wb, "Teams", ["Team", "Abbr", "Conference", "Division"], TEAMS)
    sheet(wb, "Notes", ["Date", "Note"], [
        ("09/29/2026", "Seeded from franchise screens: owner goals, budget, cap, contracts, lines, two trade offers (both declined)."),
        ("09/29/2026", "Goalie contracts (Sorokin, Varlamov) pending from the user."),
        ("09/29/2026", "Varlamov sits in goalie slot 1 on the lineup screen, Sorokin in slot 2."),
        ("09/29/2026", "Schedule loaded from the NHL.com official 2026-27 release: 84 games, 42 home, 42 away. The new CBA expands the season from 82 to 84; both added games are intra-division, so every Metropolitan rival is played 4 times. Opener is 09/30/2026 at TOR."),
    ], {"Note": 80})
    wb.save(OUT)
    print("wrote", OUT)

if __name__ == "__main__":
    main()
